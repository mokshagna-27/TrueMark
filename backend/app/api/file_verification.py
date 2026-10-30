import io
import re
import zipfile
from collections import Counter
from typing import Any

import fitz
from docx import Document
from fastapi import APIRouter, File, Form, HTTPException, Query, UploadFile
from PIL import Image, ImageChops

from app.api.core.models.services.analyzer_service import analyze_content
from app.api.router import brand_profiles


router = APIRouter(prefix="/verification", tags=["File Verification"])
MAX_UPLOAD_BYTES = 5 * 1024 * 1024

PNG_SIGNATURE = b"\x89PNG\r\n\x1a\n"
JPEG_SIGNATURES = (b"\xff\xd8\xff",)
PDF_SIGNATURE = b"%PDF-"


def finding(
    check: str,
    status: str,
    severity: str,
    explanation: str,
    suggested_fix: str,
) -> dict[str, str]:
    return {
        "check_name": check,
        "status": status,
        "severity": severity,
        "explanation": explanation,
        "suggested_fix": suggested_fix,
    }


def detect_format(filename: str, content: bytes) -> str | None:
    extension = filename.rsplit(".", 1)[-1].casefold() if "." in filename else ""

    if content.startswith(PNG_SIGNATURE) and extension == "png":
        return "png"
    if content.startswith(JPEG_SIGNATURES) and extension in {"jpg", "jpeg"}:
        return "jpg"
    if content.startswith(PDF_SIGNATURE) and extension == "pdf":
        return "pdf"
    if extension == "docx" and is_valid_docx(content):
        return "docx"
    return None


def is_valid_docx(content: bytes) -> bool:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            names = set(archive.namelist())
            return "[Content_Types].xml" in names and any(
                name.startswith("word/") for name in names
            )
    except (zipfile.BadZipFile, OSError):
        return False


def extract_pdf_text(content: bytes) -> str:
    document = fitz.open(stream=content, filetype="pdf")
    try:
        return "\n".join(page.get_text() for page in document)
    finally:
        document.close()


def extract_docx_text(content: bytes) -> str:
    document = Document(io.BytesIO(content))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            paragraphs.extend(cell.text for cell in row.cells)
    return "\n".join(paragraphs)


def hex_to_rgb(value: str) -> tuple[int, int, int] | None:
    match = re.fullmatch(r"#?([0-9a-fA-F]{6})", value.strip())
    if not match:
        return None
    return tuple(int(match.group(1)[index:index + 2], 16) for index in (0, 2, 4))


def dominant_colors(image: Image.Image, limit: int = 6) -> list[str]:
    sample = image.convert("RGBA")
    sample.thumbnail((128, 128))
    pixels = []
    for red, green, blue, alpha in sample.getdata():
        if alpha < 128:
            continue
        pixels.append((red // 32 * 32, green // 32 * 32, blue // 32 * 32))

    colors = [
        f"#{red:02x}{green:02x}{blue:02x}"
        for (red, green, blue), _count in Counter(pixels).most_common(limit)
    ]
    return colors


def color_distance(left: tuple[int, int, int], right: tuple[int, int, int]) -> float:
    return sum((a - b) ** 2 for a, b in zip(left, right)) ** 0.5


def dhash(image: Image.Image, hash_size: int = 8) -> int:
    grayscale = image.convert("L").resize(
        (hash_size + 1, hash_size), Image.Resampling.LANCZOS
    )
    pixels = list(grayscale.getdata())
    value = 0
    for row in range(hash_size):
        start = row * (hash_size + 1)
        for column in range(hash_size):
            value <<= 1
            value |= int(pixels[start + column] > pixels[start + column + 1])
    return value


def hamming_distance(left: int, right: int) -> int:
    return (left ^ right).bit_count()


@router.post("/file")
def verify_file(
    file: UploadFile = File(...),
    brand_id: str | None = Form(default=None),
    required_terms: list[str] = Query(default_factory=list),
    forbidden_terms: list[str] = Query(default_factory=list),
):
    if not file.filename:
        raise HTTPException(status_code=400, detail="A filename is required.")

    content = file.file.read(MAX_UPLOAD_BYTES + 1)
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="File exceeds the 5 MB upload limit.")

    file_format = detect_format(file.filename, content)
    if file_format is None:
        raise HTTPException(
            status_code=415,
            detail="Unsupported or invalid file. Use a genuine PNG, JPG/JPEG, PDF, or DOCX file.",
        )

    profile = brand_profiles.get(brand_id) if brand_id else None
    if brand_id and profile is None:
        raise HTTPException(status_code=404, detail="Brand profile not found.")

    if profile:
        required_terms = required_terms or profile.get("required_terms", [])
        forbidden_terms = forbidden_terms or profile.get("forbidden_terms", [])

    accessible_checks: list[dict[str, str]] = []
    unsupported_checks: list[dict[str, str]] = []

    metadata: dict[str, Any] = {
        "filename": file.filename,
        "format": file_format,
        "size_bytes": len(content),
    }

    try:
        if file_format in {"pdf", "docx"}:
            text = extract_pdf_text(content) if file_format == "pdf" else extract_docx_text(content)
            metadata["extracted_text_characters"] = len(text)
            rule_result = analyze_content(text, required_terms, forbidden_terms, channel="document")

            for rule_finding in rule_result["findings"]:
                accessible_checks.append(finding(
                    rule_finding["rule"],
                    "failed" if rule_finding.get("matched", True) else "passed",
                    rule_finding["severity"],
                    rule_finding["message"],
                    rule_finding["suggestion"],
                ))

            if not text.strip():
                accessible_checks.append(finding(
                    "readable_text",
                    "unsupported",
                    "medium",
                    "No readable text could be extracted from this document.",
                    "Provide an OCR text layer or a text-accessible copy.",
                ))
            else:
                accessible_checks.append(finding(
                    "readable_text",
                    "passed",
                    "info",
                    f"Readable text was extracted ({len(text)} characters).",
                    "No action is required.",
                ))
        else:
            with Image.open(io.BytesIO(content)) as image:
                image.verify()
            with Image.open(io.BytesIO(content)) as image:
                metadata["dimensions"] = {"width": image.width, "height": image.height}
                metadata["dominant_colors"] = dominant_colors(image)
                brand_colors = profile.get("brand_colors", []) if profile else []

                if brand_colors:
                    unmatched = []
                    configured = []
                    for value in brand_colors:
                        rgb = hex_to_rgb(value)
                        if rgb is not None:
                            configured.append((value, rgb))

                    for configured_value, configured_rgb in configured:
                        matches = any(
                            color_distance(hex_to_rgb(dominant), configured_rgb) <= 75
                            for dominant in metadata["dominant_colors"]
                            if hex_to_rgb(dominant) is not None
                        )
                        if not matches:
                            unmatched.append(configured_value)

                    if unmatched:
                        accessible_checks.append(finding(
                            "brand_colors",
                            "needs_review",
                            "medium",
                            "Configured brand colors not found in dominant colors: " + ", ".join(unmatched),
                            "Adjust image colors toward the configured brand palette or add a justified palette.",
                        ))
                    else:
                        accessible_checks.append(finding(
                            "brand_colors",
                            "passed",
                            "info",
                            "At least one dominant image color matches each valid configured brand color.",
                            "No action is required.",
                        ))
                else:
                    accessible_checks.append(finding(
                        "brand_colors",
                        "unavailable",
                        "info",
                        "Brand color verification is unavailable because the profile has no valid brand colors configured.",
                        "Add brand_colors as hexadecimal values to the brand profile.",
                    ))

                reference_logo = profile.get("reference_logo_base64") if profile else None
                if reference_logo:
                    try:
                        import base64
                        reference_content = base64.b64decode(reference_logo, validate=True)
                        with Image.open(io.BytesIO(reference_content)) as reference_image:
                            reference_hash = dhash(reference_image.convert("RGB"))
                            submitted_hash = dhash(image.convert("RGB"))
                        distance = hamming_distance(reference_hash, submitted_hash)
                        matches_logo = distance <= 10
                        accessible_checks.append(finding(
                            "reference_logo",
                            "passed" if matches_logo else "needs_review",
                            "medium" if not matches_logo else "info",
                            f"Logo comparison used dHash similarity (Hamming distance: {distance}, threshold: 10).",
                            "Use the approved logo without resizing distortion, recoloring, or background substitution." if not matches_logo else "No action is required.",
                        ))
                    except Exception:
                        accessible_checks.append(finding(
                            "reference_logo",
                            "unavailable",
                            "medium",
                            "The configured reference logo could not be decoded or compared as an image.",
                            "Store reference_logo_base64 as valid base64-encoded PNG or JPEG image data.",
                        ))
                else:
                    accessible_checks.append(finding(
                        "reference_logo",
                        "unavailable",
                        "info",
                        "Reference-logo verification is unavailable; no logo reference is configured in the brand profile.",
                        "Configure reference_logo_base64, using a clear reference logo image and the documented dHash comparison method.",
                    ))

                accessible_checks.append(finding(
                    "image_dimensions",
                    "passed",
                    "info",
                    f"Image dimensions extracted as {image.width} x {image.height} pixels.",
                    "No action is required.",
                ))
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"The file could not be verified: {exc}") from exc

    unsupported_checks.extend([
        finding("alt_text", "unsupported", "info", "Alt text is not encoded or verifiable in a PNG, JPG/JPEG, PDF, or DOCX file upload by this endpoint.", "Provide alt text in the website or document source where it is authored."),
        finding("text_contrast", "unsupported", "info", "Rendered text contrast cannot be established reliably from pixel or document bytes.", "Use accessible styling and verify contrast in the final rendered frontend or document."),
    ])

    accessible_checks = [item for item in accessible_checks if item["check_name"] != "logo_comparison_method"]
    return {
        "status": "success",
        "verification_type": "file_verification",
        "file": metadata,
        "accessible_checks": accessible_checks,
        "unsupported_checks": unsupported_checks,
        "brand_configuration_missing": {
            "brand_colors": not profile or not profile.get("brand_colors"),
            "reference_logo": not profile or not profile.get("reference_logo_base64"),
        },
    }
